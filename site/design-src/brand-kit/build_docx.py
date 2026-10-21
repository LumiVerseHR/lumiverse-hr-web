"""Word templates for the brand kit: letterhead and memo (EN, HR), A4.

Run after build-kit.mjs, which writes the header logo and kit.json:
  python build_docx.py   (needs python-docx)
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
KIT = HERE.parent.parent / "public" / "assets" / "brand" / "kit"
CO = json.loads((KIT / "kit.json").read_text())["company"]

INK = RGBColor(0x0E, 0x0F, 0x12)
SOFT = RGBColor(0x55, 0x53, 0x4F)
BASSWOOD = RGBColor(0x8A, 0x6A, 0x43)
FONT = "Golos Text"


def base():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(40), Mm(30)
    sec.header_distance, sec.footer_distance = Mm(14), Mm(10)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.35

    # Header: the lockup, flush with the text column.
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(KIT / "docx-header-logo.png"), height=Mm(8))

    # Footer: a hairline, then the legal lines every business letter carries.
    fp = sec.footer.paragraphs[0]
    rule(fp)
    for i, line in enumerate((CO["legal"], CO["capitalLine"], CO["bankLine"])):
        r = fp.add_run(("\n" if i else "") + line)
        r.font.size = Pt(7)
        r.font.color.rgb = SOFT
    return doc


def rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "6"), ("color", "D9D7D3")):
        top.set(qn(f"w:{k}"), v)
    bdr.append(top)
    pPr.append(bdr)


def letterhead():
    doc = base()
    p = doc.add_paragraph("Zagreb, ")
    p.runs[0].font.color.rgb = SOFT
    doc.add_paragraph("Primatelj\nUlica i broj\nPoštanski broj i mjesto").paragraph_format.space_before = Pt(18)
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(22)
    r = h.add_run("Predmet: ")
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph("Poštovani,").paragraph_format.space_before = Pt(10)
    doc.add_paragraph("")
    doc.add_paragraph("S poštovanjem,").paragraph_format.space_before = Pt(18)
    s = doc.add_paragraph()
    s.paragraph_format.space_before = Pt(26)
    s.add_run("Ime Prezime").bold = True
    s.add_run("\nFunkcija").font.color.rgb = SOFT
    doc.save(str(KIT / "letterhead-a4.docx"))


def memo(lang):
    L = {
        "en": ("To", "From", "Cc", "Date", "Subject"),
        "hr": ("Za", "Od", "Kopija", "Datum", "Predmet"),
    }[lang]
    doc = base()
    t = doc.add_paragraph()
    r = t.add_run("Memorandum")
    r.bold = True
    r.font.size = Pt(28)
    t.paragraph_format.space_after = Pt(14)

    tbl = doc.add_table(rows=len(L), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row, label in zip(tbl.rows, L):
        row.cells[0].width, row.cells[1].width = Mm(28), Mm(146)
        c = row.cells[0].paragraphs[0]
        lr = c.add_run(label)
        lr.italic = True
        lr.font.color.rgb = BASSWOOD
        lr.font.size = Pt(10)
        for cell in row.cells:
            borders(cell)
    doc.add_paragraph("").paragraph_format.space_before = Pt(16)
    doc.save(str(KIT / f"memo-{lang}-a4.docx"))


def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("color", "E4E2DE")):
            e.set(qn(f"w:{k}"), v)
        b.append(e)
    tcPr.append(b)


if __name__ == "__main__":
    letterhead()
    memo("en")
    memo("hr")
    for f in ("letterhead-a4.docx", "memo-en-a4.docx", "memo-hr-a4.docx"):
        print(f, (KIT / f).stat().st_size, "B")
